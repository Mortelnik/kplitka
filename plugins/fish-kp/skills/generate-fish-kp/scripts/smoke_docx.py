"""Self-contained technical probe, NOT the business calculation engine."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation, localcontext
from pathlib import Path
from uuid import uuid4


def decimal_string(data: dict, key: str) -> Decimal:
    value = data.get(key)
    if not isinstance(value, str) or not re.fullmatch(r"\d{1,12}(?:\.\d{1,6})?", value):
        raise ValueError(f"{key}: expected a nonnegative decimal string, max 6 decimals")
    return Decimal(value)


def calculate(data: dict) -> dict:
    if not isinstance(data, dict) or data.get("mode") != "technical-smoke-only":
        raise ValueError("Only mode=technical-smoke-only is implemented")
    if set(data) != {"mode", "quantity_tons", "unit_price", "tax_rate", "currency"}:
        raise ValueError("Missing or unknown input fields")
    if not isinstance(data["currency"], str) or not re.fullmatch(r"[A-Z]{3}", data["currency"]):
        raise ValueError("currency must be three uppercase ASCII letters")
    quantity = decimal_string(data, "quantity_tons")
    price = decimal_string(data, "unit_price")
    rate = decimal_string(data, "tax_rate")
    if quantity <= 0 or price <= 0 or rate > 1:
        raise ValueError("quantity and price must be positive; tax_rate must be between 0 and 1")
    with localcontext() as context:
        context.prec = 50
        cents = Decimal("0.01")
        subtotal = (quantity * price).quantize(cents, rounding=ROUND_HALF_UP)
        tax = (subtotal * rate).quantize(cents, rounding=ROUND_HALF_UP)
        total = subtotal + tax
    return {
        "quantity_tons": str(quantity),
        "unit_price": str(price),
        "tax_rate": str(rate),
        "currency": data["currency"],
        "subtotal": format(subtotal, ".2f"),
        "tax": format(tax, ".2f"),
        "total": format(total, ".2f"),
    }


def create_document(result: dict, output_dir: Path) -> Path:
    from docx import Document
    from docx.shared import Pt

    output_dir.mkdir(parents=True, exist_ok=True)
    target = output_dir / f"technical-smoke-{uuid4().hex}.docx"
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    document.add_heading("Проверка плагина fish-kp", level=0)
    document.add_paragraph(
        "Только технический пример. Все значения синтетические. "
        "Это не коммерческое предложение и не расчёт доставки."
    )
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for name, value in [
        ("Количество, т", result["quantity_tons"]),
        ("Цена за тонну", result["unit_price"]),
        ("Валюта", result["currency"]),
        ("База без налога", result["subtotal"]),
        ("Учебная ставка налога", result["tax_rate"]),
        ("Налог", result["tax"]),
        ("ИТОГО", result["total"]),
    ]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = name, value
    document.add_paragraph(
        "Расчёт выполнен Python Decimal. Здесь не проверяются контейнеры, "
        "скидки, складские остатки или налоговый режим реальных товаров."
    )
    document.save(target)
    return target.resolve()


def main() -> int:
    # A Cyrillic output path must survive a redirected Windows console.
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args()
    try:
        data = json.loads(args.input.read_text(encoding="utf-8"))
        result = calculate(data)
        target = create_document(result, args.output_dir)
        payload = {
            "ok": True,
            "mode": "technical-smoke-only",
            "calculation": result,
            "document_path": str(target),
            "sha256": hashlib.sha256(target.read_bytes()).hexdigest(),
            "client_delivery_verified": False,
        }
    except (ValueError, OSError, ImportError, InvalidOperation) as error:
        print(json.dumps({"ok": False, "error": str(error)}, ensure_ascii=False))
        return 2
    print(json.dumps(payload, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())

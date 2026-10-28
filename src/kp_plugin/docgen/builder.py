"""Low-level DOCX constructor without commercial approval checks.

Use docgen.gate.generate_commercial_document for the checked commercial path.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from kp_plugin.docgen.corporate import (
    CorporateOffer,
    apply_corporate_style,
    build_corporate_document,
)

REQUIRED_FIELDS = (
    "product",
    "product_source",
    "price_source",
    "containers",
    "pickup_tons",
    "base_amount",
    "discount_rule",
    "discount_amount",
    "tax",
    "total",
)


@dataclass
class OfferPosition:
    product: str
    quantity: Any
    unit: str
    unit_price: Any | None = None
    line_total: Any | None = None
    note: str | None = None

    def validate(self) -> None:
        if not isinstance(self.product, str) or not self.product.strip():
            raise ValueError("Обязательное поле: product")
        if self.quantity is None:
            raise ValueError("Обязательное поле: quantity")
        if not isinstance(self.unit, str) or not self.unit.strip():
            raise ValueError("Обязательное поле: unit")
        if self.line_total is None:
            raise ValueError("Обязательное поле: line_total")

        self.quantity = _as_decimal(self.quantity, "quantity")
        self.line_total = _as_decimal(self.line_total, "line_total")
        if self.unit_price is not None:
            self.unit_price = _as_decimal(self.unit_price, "unit_price")

    def as_display_row(self) -> dict[str, str]:
        self.validate()
        return {
            "product": self.product.strip(),
            "quantity": _format_quantity(self.quantity),
            "unit": self.unit,
            "unit_price": _format_money(self.unit_price) if self.unit_price is not None else "—",
            "line_total": _format_money(self.line_total),
        }


@dataclass
class OfferPresentation:
    document_date: str
    client_task: str
    proposed_solution: str
    items: list[OfferPosition]
    totals: dict[str, Any] = field(default_factory=dict)
    seller: dict[str, str] | None = None
    buyer: dict[str, str] | None = None
    delivery_terms: list[str] = field(default_factory=list)
    contacts: list[str] = field(default_factory=list)
    technical_summary: dict[str, Any] | None = None
    is_test_data: bool = False
    source_note: str | None = None

    def validate(self) -> None:
        if not isinstance(self.document_date, str) or not self.document_date.strip():
            raise ValueError("Обязательное поле: document_date")
        if not isinstance(self.client_task, str) or not self.client_task.strip():
            raise ValueError("Обязательное поле: client_task")
        if not isinstance(self.proposed_solution, str) or not self.proposed_solution.strip():
            raise ValueError("Обязательное поле: proposed_solution")
        if not isinstance(self.items, list) or not self.items:
            raise ValueError("Обязательное поле: items")
        for item in self.items:
            if not isinstance(item, OfferPosition):
                raise TypeError("Обязательное поле: items")
            item.validate()

        missing = [field for field in ("base_amount", "discount_amount", "tax", "total") if field not in self.totals or self.totals[field] is None]
        if missing:
            raise ValueError(f"Обязательное поле: totals[{missing[0]}]")

        for field_name in ("base_amount", "discount_amount", "tax", "total"):
            self.totals[field_name] = _as_decimal(self.totals[field_name], field_name)

        line_total_sum = sum((item.line_total for item in self.items), Decimal("0.00"))
        if line_total_sum != self.totals["base_amount"]:
            raise ValueError(
                "Несовпадение суммы строк item.line_total с totals['base_amount']: "
                f"{line_total_sum} != {self.totals['base_amount']}"
            )

        for value in (self.document_date, self.client_task, self.proposed_solution):
            _ensure_customer_safe_text(value, "customer text")

        if self.is_test_data:
            self.source_note = self.source_note or "Демонстрационный сценарий: тестовые данные, расчёт P2 не подтверждён."


def _as_decimal(value: Any, field_name: str) -> Decimal:
    if isinstance(value, Decimal):
        return value
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        return Decimal(str(value))
    if isinstance(value, str):
        text = value.strip().replace(" ", "")
        if not text:
            raise ValueError(f"Отсутствует обязательное поле: {field_name}")
        try:
            return Decimal(text)
        except InvalidOperation as exc:  # pragma: no cover - defensive guard
            raise ValueError(f"Некорректное денежное значение для поля {field_name}: {value!r}") from exc
    raise ValueError(f"Отсутствует обязательное поле: {field_name}")


def _format_money(value: Any) -> str:
    decimal_value = _as_decimal(value, "money")
    formatted = f"{decimal_value.quantize(Decimal('0.01')):,.2f}"
    formatted = formatted.replace(",", " ")
    formatted = formatted.replace(".", ",")
    return formatted + " ₽"


def _format_quantity(value: Any) -> str:
    decimal_value = _as_decimal(value, "quantity")
    formatted = f"{decimal_value.quantize(Decimal('0.01')):,.2f}"
    formatted = formatted.replace(",", " ")
    formatted = formatted.replace(".", ",")
    return formatted


def _ensure_customer_safe_text(value: Any, field_name: str) -> None:
    text = str(value)
    if ".csv" in text.lower() or "product_source" in text or "price_source" in text or "source_ids" in text:
        raise ValueError(f"Технические идентификаторы не допускаются в клиентском документе: {field_name}")


def _normalize_offer_result(result: Any) -> dict[str, Any]:
    if isinstance(result, str):
        raise TypeError(f"Ошибка расчёта: {result}")
    if not isinstance(result, dict):
        raise TypeError("Ошибка расчёта: результат P2 должен быть словарём")

    missing = [field for field in REQUIRED_FIELDS if field not in result or result[field] is None]
    if missing:
        raise ValueError(f"Отсутствует обязательное поле: {missing[0]}")

    normalized = dict(result)
    for total_field in ("base_amount", "discount_amount", "tax", "total"):
        normalized[total_field] = _as_decimal(result[total_field], total_field)
    if isinstance(result.get("pickup_tons"), (int, float, str, Decimal)):
        normalized["pickup_tons"] = _as_decimal(result["pickup_tons"], "pickup_tons")
    return normalized


def _build_legacy_offer_document(offer: dict[str, Any], path: Path) -> None:
    document = Document()
    apply_corporate_style(document)
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(10)

    document.add_heading("Коммерческое предложение", level=0)
    document.add_paragraph(f"Продукт: {offer['product']}")

    document.add_heading("Расчёт", level=1)
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    headers = table.rows[0].cells
    headers[0].text = "Параметр"
    headers[1].text = "Значение"

    rows = [
        ("Контейнеры", str(offer["containers"])),
        ("Подборка, т", _format_tons(offer["pickup_tons"])),
        ("Базовая сумма", _format_money(offer["base_amount"])),
        ("Правило скидки", str(offer["discount_rule"])),
        ("Скидка", _format_money(offer["discount_amount"])),
        ("НДС", _format_money(offer["tax"])),
        ("ИТОГО", _format_money(offer["total"])),
    ]
    for label, value in rows:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = value

    document.add_heading("Источники", level=1)
    document.add_paragraph(f"Продукт: {offer['product_source']}")
    document.add_paragraph(f"Прайс: {offer['price_source']}")
    document.add_paragraph(f"Скидка: {offer['discount_rule']}")

    document.save(path)


def _build_offer_presentation_document(presentation: OfferPresentation, path: Path) -> None:
    presentation.validate()

    document = Document()
    apply_corporate_style(document, (presentation.seller or {}).get("name"))
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(10)

    document.add_heading("Коммерческое предложение", level=0)
    if presentation.is_test_data and presentation.source_note:
        document.add_paragraph("ТЕСТОВЫЙ ОБРАЗЕЦ — НЕ ДЛЯ КОММЕРЧЕСКОГО ИСПОЛЬЗОВАНИЯ (TEST ONLY)")
        note = document.add_paragraph()
        note.alignment = WD_ALIGN_PARAGRAPH.LEFT
        note.add_run(presentation.source_note)

    if presentation.document_date:
        document.add_paragraph(f"Дата: {_format_russian_date(presentation.document_date)}")

    if presentation.seller and presentation.seller.get("name"):
        seller_name = presentation.seller.get("name", "").strip()
        if seller_name:
            document.add_paragraph(f"Продавец: {seller_name}")
    if presentation.buyer and presentation.buyer.get("name"):
        buyer_name = presentation.buyer.get("name", "").strip()
        if buyer_name:
            document.add_paragraph(f"Покупатель: {buyer_name}")

    if presentation.contacts:
        document.add_heading("Контакты", level=1)
        for contact in presentation.contacts:
            _ensure_customer_safe_text(contact, "contacts")
            document.add_paragraph(contact)

    document.add_heading("Задача клиента", level=1)
    document.add_paragraph(presentation.client_task)

    document.add_heading("Предлагаемое решение", level=1)
    document.add_paragraph(presentation.proposed_solution)

    document.add_heading("Состав поставки", level=1)
    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_ALIGN_PARAGRAPH.LEFT
    table.autofit = False
    table.columns[0].width = Inches(2.8)
    table.columns[1].width = Inches(0.95)
    table.columns[2].width = Inches(0.7)
    table.columns[3].width = Inches(1.2)
    table.columns[4].width = Inches(1.35)

    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(8)

    headers = table.rows[0].cells
    headers[0].text = "Товар"
    headers[1].text = "Количество"
    headers[2].text = "Ед. изм."
    headers[3].text = "Цена за единицу"
    headers[4].text = "Сумма"
    for cell in headers:
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(8)

    for item in presentation.items:
        row = item.as_display_row()
        cells = table.add_row().cells
        cells[0].text = row["product"]
        cells[1].text = row["quantity"]
        cells[2].text = row["unit"]
        cells[3].text = row["unit_price"]
        cells[4].text = row["line_total"]
        for cell in cells:
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(8)

    document.add_heading("Итоговые суммы", level=1)
    totals_table = document.add_table(rows=1, cols=2)
    totals_table.style = "Table Grid"
    totals_table.columns[0].width = 3500
    totals_table.columns[1].width = 2200
    totals_headers = totals_table.rows[0].cells
    totals_headers[0].text = "Параметр"
    totals_headers[1].text = "Значение"

    totals_rows = [
        ("Базовая сумма", _format_money(presentation.totals["base_amount"])),
        ("Скидка", _format_money(presentation.totals["discount_amount"])),
        ("НДС", _format_money(presentation.totals["tax"])),
        ("Итого к оплате", _format_money(presentation.totals["total"])),
    ]
    for label, value in totals_rows:
        row = totals_table.add_row().cells
        row[0].text = label
        row[1].text = value
        if label == "Итого к оплате":
            for paragraph in row[0].paragraphs:
                for run in paragraph.runs:
                    run.bold = True
            for paragraph in row[1].paragraphs:
                for run in paragraph.runs:
                    run.bold = True

    if presentation.delivery_terms:
        document.add_heading("Условия поставки", level=1)
        for term in presentation.delivery_terms:
            _ensure_customer_safe_text(term, "delivery_terms")
            paragraph = document.add_paragraph(term)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.widow_control = True

    document.save(path)


def _format_russian_date(value: str) -> str:
    if not value:
        return ""
    try:
        year, month, day = value.split("-")
    except ValueError:
        return value
    month_names = {
        "01": "января",
        "02": "февраля",
        "03": "марта",
        "04": "апреля",
        "05": "мая",
        "06": "июня",
        "07": "июля",
        "08": "августа",
        "09": "сентября",
        "10": "октября",
        "11": "ноября",
        "12": "декабря",
    }
    return f"{int(day)} {month_names.get(month, month)} {int(year)} г."


def build_offer_document(result: Any, output_path: str | Path) -> Path:
    """Create a readable Word document from a P2 calculation result or P3 offer preview."""
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    if isinstance(result, CorporateOffer):
        build_corporate_document(result, path)
        return path

    if isinstance(result, OfferPresentation):
        _build_offer_presentation_document(result, path)
        return path

    offer = _normalize_offer_result(result)
    _build_legacy_offer_document(offer, path)
    return path


def _format_tons(value: Any) -> str:
    decimal_value = _as_decimal(value, "pickup_tons")
    formatted = f"{decimal_value.quantize(Decimal('0.01')):,.2f}"
    formatted = formatted.replace(",", " ")
    formatted = formatted.replace(".", ",")
    return formatted + " т"

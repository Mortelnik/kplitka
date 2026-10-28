"""Corporate Word layout only. All business values must be supplied by the checked caller."""

from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

NAVY = "163A59"
PALE = "EAF2F6"
HEADERS = ("Наименование", "Количество", "Цена за тонну", "Скидка", "НДС", "Итого")
WIDTHS = (51, 23, 28, 23, 22, 23)  # 170 mm = A4 minus two 20 mm margins.


@dataclass(frozen=True)
class DemoCompanyDetails:
    """Fictional presentation content: accepted by gate ONLY in explicit TEST ONLY mode."""

    company_name: str
    business_description: str
    client_name: str
    delivery_terms: tuple[str, ...]
    contacts: tuple[str, ...]


def russian_date(value: str) -> str:
    parsed = date.fromisoformat(value)
    months = ("января", "февраля", "марта", "апреля", "мая", "июня", "июля", "августа",
              "сентября", "октября", "ноября", "декабря")
    return f"{parsed.day} {months[parsed.month - 1]} {parsed.year} г."


@dataclass(frozen=True)
class CorporateOffer:
    product: str
    quantity: Decimal
    unit: str
    totals: dict[str, Decimal]
    currency: str
    document_date: str
    unit_price: Decimal | None = None
    company_name: str | None = None
    client_name: str | None = None
    offer_number: str | None = None
    business_description: str | None = None
    subtitle: str | None = None
    delivery_terms: tuple[str, ...] = ()
    notes: tuple[str, ...] = ()
    contacts: tuple[str, ...] = ()
    cases: tuple[str, ...] = ()
    test_only: bool = False
    internal_draft: bool = False
    # Provenance stays in the envelope, not in customer-facing paragraphs.
    provenance: dict[str, str] = field(default_factory=dict)


def _shade(cell, fill):
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(element)


def apply_corporate_style(document, company_name=None):
    """Real Word styles, A4 geometry and a PAGE field; no made-up company identity."""
    for section in document.sections:
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin = section.bottom_margin = Mm(20)
        section.left_margin = section.right_margin = Mm(20)
        section.header_distance = section.footer_distance = Mm(10)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if company_name:
            footer.add_run(company_name + "  •  ")
        footer.add_run("Страница ")
        page = OxmlElement("w:fldSimple")
        page.set(qn("w:instr"), "PAGE")
        footer._p.append(page)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.widow_control = True
    normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
    for name, size in (("Title", 23), ("Heading 1", 12), ("Heading 2", 11)):
        style = document.styles[name]
        style.font.name, style.font.size, style.font.color.rgb = "Arial", Pt(size), RGBColor.from_string(NAVY)
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
        if name != "Title":
            style.font.bold = True
    company_style = document.styles.add_style("Company Name", WD_STYLE_TYPE.PARAGRAPH)
    company_style.base_style = document.styles["Normal"]
    company_style.font.name = "Arial"
    company_style.font.size = Pt(14)
    company_style.font.bold = True
    company_style.font.color.rgb = RGBColor.from_string(NAVY)
    company_style.paragraph_format.space_before = Pt(10)
    company_style.paragraph_format.space_after = Pt(6)
    company_style.paragraph_format.keep_with_next = True


def _money(value, currency):
    if not isinstance(value, Decimal) or not value.is_finite():
        raise ValueError("A finite original Decimal is required")
    symbol = "₽" if currency == "RUB" else currency
    return f"{value:,.2f}".replace(",", " ").replace(".", ",") + " " + symbol


def _table_layout(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for column, width in zip(table.columns, widths, strict=True):
        column.width = Mm(width)
    for row in table.rows:
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for index, (cell, width) in enumerate(zip(row.cells, widths, strict=True)):
            cell.width = Mm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ("top", "bottom", "left", "right"):
                node = OxmlElement("w:" + side)
                node.set(qn("w:w"), "100")
                node.set(qn("w:type"), "dxa")
                margins.append(node)
            props.append(margins)
            borders = OxmlElement("w:tcBorders")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "4")
            bottom.set(qn("w:color"), "D7E2EA")
            borders.append(bottom)
            props.append(borders)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.RIGHT
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.keep_together = True
                for run in paragraph.runs:
                    run.font.size = Pt(9)


def build_corporate_document(offer: CorporateOffer, path: Path):
    """Low-level renderer; commercial callers must use gate.generate_commercial_document."""
    document = Document()
    apply_corporate_style(document, offer.company_name)
    if offer.internal_draft:
        status = "INTERNAL_DRAFT / TEST ONLY — НЕ ДЛЯ КЛИЕНТА"
        paragraph = document.add_paragraph(status)
        paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            run.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor.from_string(NAVY)
        for section in document.sections:
            section.footer.add_paragraph(status)
    if offer.test_only:
        banner = document.add_paragraph(
            "ТЕСТОВЫЙ ОБРАЗЕЦ — НЕ ДЛЯ КОММЕРЧЕСКОГО ИСПОЛЬЗОВАНИЯ (TEST ONLY)")
        banner.paragraph_format.keep_with_next = True
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), PALE)
        banner._p.get_or_add_pPr().append(shade)
        for run in banner.runs:
            run.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor.from_string(NAVY)
    if offer.company_name:
        document.add_paragraph(offer.company_name, "Company Name")
    if offer.business_description:
        document.add_paragraph(offer.business_description)
    if offer.offer_number:
        document.add_paragraph(f"КП № {offer.offer_number}")
    document.add_heading("Коммерческое предложение", 0)
    if offer.subtitle:
        document.add_paragraph(offer.subtitle)
    if offer.client_name:
        document.add_paragraph(f"Клиент: {offer.client_name}")
    document.add_paragraph(f"Дата: {russian_date(offer.document_date)}")
    document.add_heading("Спецификация и стоимость", 1)
    table = document.add_table(rows=2, cols=6)
    for cell, text in zip(table.rows[0].cells, HEADERS, strict=True):
        cell.text = text
        _shade(cell, NAVY)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    values = (offer.product, f"{offer.quantity} {offer.unit}",
              _money(offer.unit_price, offer.currency) if offer.unit_price is not None and offer.unit == "т" else "Не подтверждена",
              _money(offer.totals["discount_amount"], offer.currency), _money(offer.totals["tax"], offer.currency),
              _money(offer.totals["total"], offer.currency))
    for cell, text in zip(table.rows[1].cells, values, strict=True):
        cell.text = text
        _shade(cell, "F7FAFC")
    _table_layout(table, WIDTHS)
    totals = document.add_table(rows=4, cols=2)
    labels = ("Базовая стоимость", "Скидка", "НДС", "ОБЩАЯ СУММА С НДС")
    for row, label, key in zip(totals.rows, labels, ("base_amount", "discount_amount", "tax", "total"), strict=True):
        row.cells[0].text, row.cells[1].text = label, _money(offer.totals[key], offer.currency)
    _table_layout(totals, (115, 55))
    for cell in totals.rows[-1].cells:
        _shade(cell, PALE)
        for run in cell.paragraphs[0].runs:
            run.bold, run.font.size, run.font.color.rgb = True, Pt(13), RGBColor.from_string(NAVY)
    first_section = True
    for title, lines in (("Условия поставки", offer.delivery_terms), ("Примечание", offer.notes),
                         ("Релевантные кейсы", offer.cases), ("Контакты", offer.contacts)):
        if lines:
            heading = document.add_heading(title, 1)
            if first_section or title == "Контакты":
                heading.paragraph_format.space_before = Pt(19)  # 5 pt extra after totals / before contacts.
            first_section = False
            for text in lines:
                document.add_paragraph(text)
    document.save(path)

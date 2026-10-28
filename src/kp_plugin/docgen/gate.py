"""Fail-closed local commercial gate; attestations come from a trusted caller, not hashes."""

from copy import deepcopy
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from enum import StrEnum
from pathlib import Path

from docx import Document

from kp_plugin.docgen.builder import build_offer_document
from kp_plugin.docgen.corporate import (
    CorporateOffer,
    DemoCompanyDetails,
    apply_corporate_style,
)
from kp_plugin.llm.cases import CaseRecord, CaseStatus, CommercialPermission
from kp_plugin.llm.p2_adapter import (
    P2AdapterResult,
    P2Invocation,
    prepare_p2_invocation,
)
from kp_plugin.llm.selection import SelectionStatus, select_products
from kp_plugin.llm.selection_bridge import (
    SelectionAssessment,
    assess_selection,
    canonical_fingerprint,
)
from kp_plugin.llm.unreachable import ConfirmedSource


class GateStatus(StrEnum):
    INTERNAL_DRAFT = "INTERNAL_DRAFT"
    COMMERCIALLY_APPROVED = "COMMERCIALLY_APPROVED"


class ConstraintStatus(StrEnum):
    UNKNOWN = "unknown"
    CONFIRMED = "confirmed"
    CONTRADICTORY = "contradictory"


@dataclass(frozen=True)
class ConstraintReview:
    """External structured verification, never interpretation of customer prose.

    context_fingerprint binds the full invocation and output, including requirements,
    product, quantities, prices/rules, delivery/terms and their versions/evidence.
    Production must authenticate the reviewer and the supplied sources.
    """

    status: ConstraintStatus = ConstraintStatus.UNKNOWN
    context_fingerprint: str | None = None
    sources: tuple[ConfirmedSource, ...] = ()
    absence_confirmed: bool = False
    product_compatible: bool = False
    calculation_compatible: bool = False
    terms_compatible: bool = False


def constraint_review_fingerprint(result: P2AdapterResult, current: P2Invocation) -> str:
    return canonical_fingerprint({"invocation": current.fingerprint, "output": result.raw_output})


@dataclass(frozen=True)
class CommercialAttestation:
    """Explicit external confirmation, bound to inputs AND output; no default trust.

    Production callers must authorize their issuer. This local module cannot authenticate issuers.
    """

    issuer: str
    invocation_fingerprint: str
    output_fingerprint: str
    valid_until: date
    price_sources: tuple[ConfirmedSource, ...]
    rule_sources: tuple[ConfirmedSource, ...]
    scope_complete: bool = False
    constraints: ConstraintReview | None = None


@dataclass(frozen=True)
class GateDecision:
    status: GateStatus
    blockers: tuple[str, ...]
    warnings: tuple[str, ...]
    assessment: SelectionAssessment
    calculation: P2AdapterResult
    commercial_cases: tuple[CaseRecord, ...]


def _constraint_issues(result, current, assessment, attestation):
    issues = []
    subjects = [("request", assessment.needs),
                *((f"items[{i}]", item) for i, item in enumerate(assessment.needs.items))]
    for path, subject in subjects:
        if "constraints" in subject.unknown_fields:
            issues.append(f"constraints_unknown:{path}")
        if any(not evidence.verified for key, evidence in subject.evidence.items()
               if key == "constraints" or key.startswith("constraint:")):
            issues.append(f"constraints_unconfirmed:{path}")
    has_constraints = any(subject.constraints for _, subject in subjects)
    review = attestation.constraints if attestation else None
    if review is None or review.status == ConstraintStatus.UNKNOWN:
        issues.append("constraints_review_missing_or_unknown")
        return issues
    if review.status != ConstraintStatus.CONFIRMED:
        issues.append("constraints_contradictory")
    if not review.sources or any(not isinstance(s, ConfirmedSource) for s in review.sources):
        issues.append("constraints_sources_unconfirmed")
    if current is None or review.context_fingerprint != constraint_review_fingerprint(result, current):
        issues.append("constraints_review_stale")
    if has_constraints and review.absence_confirmed:
        issues.append("constraints_presence_conflict")
    if not has_constraints and not review.absence_confirmed:
        issues.append("constraints_absence_unconfirmed")
    for name in ("product_compatible", "calculation_compatible", "terms_compatible"):
        if getattr(review, name) is not True:
            issues.append(f"constraints_incompatible:{name}")
    return issues


def _review_covers_constraint_only_problem(problem):
    return (problem.request.reason.value == "missing_data" and bool(problem.request.missing_data)
            and all(d.field_path.endswith(".request.constraints") for d in problem.request.missing_data))


def _binding_with_review(binding, assessment):
    if binding.confirmed:
        return True
    position = next(p for p in assessment.selection.positions if p.request_position == binding.request_position)
    if len(position.candidates) != 1 or position.status != SelectionStatus.NEEDS_REVIEW:
        return False
    candidate = position.candidates[0]
    approved = {(s.source_id, s.version, s.locator) for s in assessment.confirmed_sources}
    return (candidate.product_id == binding.product_id and candidate.status == SelectionStatus.NEEDS_REVIEW
            and set(candidate.unknown_data) == {"request.constraints"} and not candidate.exclusion_reasons
            and bool(candidate.sources) and all((s.source_id, s.version, s.locator) in approved for s in candidate.sources))


def evaluate_gate(result: P2AdapterResult, current: P2Invocation | None, assessment: SelectionAssessment, *,
                  attestation: CommercialAttestation | None = None, as_of: date | None = None) -> GateDecision:
    """Rebuild checks from detached current data; never accept a pre-issued approval token."""
    result, current, assessment = deepcopy((result, current, assessment))
    constraint_issues = _constraint_issues(result, current, assessment, attestation)
    if current is None:
        return GateDecision(GateStatus.INTERNAL_DRAFT, ("calculation_context_missing", *result.diagnostics, *constraint_issues),
                            (str(assessment), str(result)), assessment, result, ())
    blockers, warnings = list(constraint_issues), []
    today = date.today() if as_of is None else as_of
    if len(assessment.needs.items) != 1:
        blockers.append("unsupported_multi_item")
    ids = {p.request_position: p.candidates[0].product_id for p in assessment.selection.positions
           if len(p.candidates) == 1}
    if assessment.selection == select_products(assessment.needs, assessment.profile.products):
        ids = None
    try:
        assessment.needs.validate()
        rebuilt = prepare_p2_invocation(assessment.needs, assessment.selection, assessment.profile,
                                        catalog=current.catalog, prices=current.prices, rules=current.rules,
                                        settings=current.settings, product_ids=ids)
        if rebuilt.fingerprint != current.fingerprint:
            blockers.append("current_context_mismatch")
        assessment = assess_selection(assessment.needs, assessment.selection, assessment.profile,
                                      product_ids=ids, confirmed_sources=assessment.confirmed_sources,
                                      unavailable=assessment.unavailability_evidence,
                                      calculation=result.evidence_for(rebuilt), calculation_context=rebuilt.bridge_context,
                                      budget=assessment.budget, case_records=assessment.case_records,
                                      client_task=assessment.client_task, task_sources=assessment.task_sources)
    except (ValueError, TypeError, ArithmeticError, IndexError, KeyError) as exc:
        blockers.append(f"invalid_current_data:{exc}")
    if not result.successful_calculation or result.diagnostics:
        blockers.append("unsuccessful_calculation")
    if not result.trusted_inputs or not current.settings.trusted_inputs:
        blockers.append("untrusted_inputs")
    evidence = result.evidence_for(current)
    if evidence is None or assessment.calculation != evidence:
        blockers.append("stale_or_missing_calculation")
    if not result.complete_delivery_cost or current.base_context.unknown_components:
        blockers.append("incomplete_cost")
    if evidence is not None:
        if evidence.cost.unknown_components or evidence.cost.delivery_cost is None:
            blockers.append("unknown_costs")
        if not isinstance(result.raw_output, dict) or result.raw_output.get("total") != evidence.cost.known_total:
            blockers.append("output_total_mismatch")
    unresolved_problems = [p for p in assessment.problems
                           if constraint_issues or not _review_covers_constraint_only_problem(p)]
    if unresolved_problems or assessment.calculation_issues:
        blockers.append("blocking_problems")
    if len(assessment.bindings) != len(assessment.needs.items) or not all(
        b.confirmed or (not constraint_issues and _binding_with_review(b, assessment)) for b in assessment.bindings
    ):
        blockers.append("unconfirmed_product_bindings")
    for position in assessment.selection.positions:
        for candidate in position.candidates:
            if candidate.requirement_changes:
                blockers.append("unapproved_requirement_changes")
    if any(row.issues or dict(row.fields).get("is_fictional") != "нет"
           for row in (*assessment.profile.catalog_rows, *assessment.profile.stock_rows)):
        blockers.append("unverified_p4_data")
    if any(row.get("is_fictional") != "нет" for row in (*current.catalog.values(), *current.prices.values())):
        blockers.append("synthetic_or_unknown_origin")
    if attestation is None:
        blockers.append("commercial_attestation_missing")
    else:
        if (not attestation.issuer.strip() or not attestation.scope_complete or today > attestation.valid_until
                or attestation.invocation_fingerprint != current.fingerprint
                or attestation.output_fingerprint != canonical_fingerprint(result.raw_output)):
            blockers.append("invalid_or_expired_attestation")
        if not attestation.price_sources or not attestation.rule_sources:
            blockers.append("price_or_rule_confirmation_missing")
        if any(not s.version.strip() or not s.confirmed for s in (*attestation.price_sources, *attestation.rule_sources)):
            blockers.append("invalid_attestation_sources")
        price_ids = {s.source_id for s in attestation.price_sources}
        if any(row["source_ref"] not in price_ids for row in current.prices.values()):
            blockers.append("price_source_unconfirmed")
    safe_cases = []
    for match in assessment.case_matches:
        record = next((r for r in assessment.case_records if r.case_id == match.case_id), None)
        if (match in assessment.commercial_cases and match.commercial_use_allowed and record is not None
                and record.status == CaseStatus.VERIFIED and record.source_confirmed and record.content_confirmed
                and record.content and record.sources and not record.is_fictional
                and record.commercial_permission == CommercialPermission.ALLOWED
                and not record.permission_conflict and (not record.usage_limits or record.usage_limits_reviewed)):
            safe_cases.append(record)
        else:
            warnings.append(f"case_omitted:{match.case_id}:{match.reasons}")
    warnings.extend(result.diagnostics)
    warnings.extend(str(p) for p in assessment.problems)
    warnings.extend(assessment.calculation_issues)
    warnings.extend(str(p) for p in assessment.selection.positions)
    blockers = tuple(dict.fromkeys(blockers))
    return GateDecision(GateStatus.INTERNAL_DRAFT if blockers else GateStatus.COMMERCIALLY_APPROVED,
                        blockers, tuple(warnings), assessment, result, tuple(safe_cases))


def generate_commercial_document(result: P2AdapterResult, current: P2Invocation,
                                 assessment: SelectionAssessment, output_path: str | Path, *,
                                 attestation: CommercialAttestation | None = None,
                                 as_of: date | None = None, test_only: bool = False,
                                 demo_details: DemoCompanyDetails | None = None) -> Path:
    """Only commercial entry point: re-evaluate every call; accepts no approval bypass flag."""
    if demo_details is not None and not test_only:
        raise ValueError("Fictional company details require explicit TEST ONLY mode")
    decision = evaluate_gate(result, current, assessment, attestation=attestation, as_of=as_of)
    if decision.status != GateStatus.COMMERCIALLY_APPROVED:
        raise ValueError(f"Commercial generation blocked: {decision.blockers}")
    invocation = decision.calculation.invocation
    binding = decision.assessment.bindings[0]
    item = decision.assessment.needs.items[binding.request_position - 1]
    offer = CorporateOffer(
        product=decision.calculation.raw_output["product"], quantity=invocation.quantity_tons, unit=item.unit,
        totals={key: decision.calculation.raw_output[key] for key in ("base_amount", "discount_amount", "tax", "total")},
        currency=invocation.settings.currency, document_date=(as_of or date.today()).isoformat(),
        unit_price=Decimal(invocation.prices[binding.product_id]["unit_price_rub_per_t"]),
        cases=tuple(record.content for record in decision.commercial_cases),
        provenance={"product_id": binding.product_id, "price_snapshot_id": invocation.settings.price_snapshot_id},
        test_only=test_only,
        company_name=demo_details.company_name if demo_details else None,
        business_description=demo_details.business_description if demo_details else None,
        client_name=demo_details.client_name if demo_details else None,
        delivery_terms=demo_details.delivery_terms if demo_details else (),
        contacts=demo_details.contacts if demo_details else (),
    )
    path = build_offer_document(offer, output_path)
    return path


def generate_internal_draft(decision: GateDecision, output_path: str | Path) -> Path:
    """Always marked INTERNAL_DRAFT; retain complete diagnostic snapshots, even without prices."""
    document = Document()
    apply_corporate_style(document)
    document.add_heading("ВНУТРЕННИЙ ЧЕРНОВИК — НЕ ДЛЯ КЛИЕНТА", 0)
    document.add_paragraph("INTERNAL_DRAFT")
    document.add_paragraph("Коммерческий допуск не предоставлен. Стоимость не подтверждена.")
    document.add_heading("Необходимые проверки и уточнения", 1)
    for blocker in decision.blockers:
        document.add_paragraph(blocker)
    document.add_heading("Внутренняя техническая диагностика", 1)
    for value in (decision.blockers, decision.warnings, decision.assessment, decision.calculation):
        document.add_paragraph(str(value))
    path = Path(output_path)
    document.save(path)
    return path


def generate_priced_internal_draft(decision: GateDecision, output_path: str | Path, *,
                                  as_of: date) -> Path:
    """Compact TEST ONLY review, never commercial approval or a substitute for diagnostics.

    No company identity, contacts or delivery promises are inferred. The caller keeps
    the full GateDecision alongside this presentation. Amounts remain the original
    successful P2 output, explicitly unconfirmed for commercial use.
    """
    if decision.status != GateStatus.INTERNAL_DRAFT:
        raise ValueError("This review path requires an INTERNAL_DRAFT decision")
    result = decision.calculation
    if not result.successful_calculation or result.invocation is None or not isinstance(result.raw_output, dict):
        raise ValueError("A successful calculation is required for a priced internal review")
    invocation = result.invocation
    offer = CorporateOffer(
        product=result.raw_output["product"], quantity=invocation.quantity_tons, unit="т",
        totals={key: result.raw_output[key] for key in ("base_amount", "discount_amount", "tax", "total")},
        currency=invocation.settings.currency, document_date=as_of.isoformat(),
        unit_price=Decimal(invocation.prices[invocation.product_id]["unit_price_rub_per_t"]),
        internal_draft=True,
        subtitle="Расчётный макет. Товар и цены тестовые; коммерческая достоверность не подтверждена.",
        delivery_terms=tuple(f"{field}: Уточняется при согласовании заказа"
                             for field in ("Способ поставки", "Срок поставки", "Оплата")),
        contacts=("Контактные данные не предоставлены.",),
        notes=("Стоимость приведена только для проверки оформления, не является коммерческим предложением.",),
    )
    return build_offer_document(offer, output_path)

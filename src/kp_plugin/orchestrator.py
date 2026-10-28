"""Local synchronous composition. No LLM/network, prices or attestations are invented here."""

from collections.abc import Callable
from copy import deepcopy
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from docx import Document

from kp_plugin.docgen.corporate import DemoCompanyDetails
from kp_plugin.docgen.gate import (
    CommercialAttestation,
    GateDecision,
    GateStatus,
    evaluate_gate,
    generate_commercial_document,
    generate_internal_draft,
)
from kp_plugin.llm.cases import CaseRecord
from kp_plugin.llm.clarify import (
    Clarification,
    ClarificationUpdateResult,
    apply_clarification_answer,
    build_clarifications,
)
from kp_plugin.llm.needs import CustomerNeeds, NeedsExtractor
from kp_plugin.llm.p2_adapter import (
    P2AdapterResult,
    P2Settings,
    prepare_p2_invocation,
    run_p2,
)
from kp_plugin.llm.p4_catalog import P4Catalog
from kp_plugin.llm.selection import SelectionResult, select_products
from kp_plugin.llm.selection_bridge import SelectionAssessment, assess_selection
from kp_plugin.llm.unreachable import ConfirmedSource


@dataclass(frozen=True)
class Answer:
    field_name: str
    position: int | None
    value: object


@dataclass(frozen=True)
class FlowResult:
    raw_text: str
    extracted: CustomerNeeds
    needs: CustomerNeeds
    initial_questions: tuple[Clarification, ...]
    updates: tuple[ClarificationUpdateResult, ...]
    remaining_questions: tuple[Clarification, ...]
    selection: SelectionResult
    calculation: P2AdapterResult
    assessment: SelectionAssessment
    decision: GateDecision
    document: Path


def run_flow(raw_text: str, *, extractor: NeedsExtractor, profile: P4Catalog, settings: P2Settings,
             catalog: dict, prices: dict, rules: dict, calculate_offer: Callable[..., object],
             output_path: str | Path, answers: tuple[Answer, ...] = (),
             product_ids: dict[int, str] | None = None, confirmed_sources: tuple[ConfirmedSource, ...] = (),
             case_records: tuple[CaseRecord, ...] = (), client_task: str | None = None,
             task_sources: tuple[ConfirmedSource, ...] = (), attestation: CommercialAttestation | None = None,
             attest: Callable[[P2AdapterResult], CommercialAttestation] | None = None,
             as_of: date | None = None, test_only: bool = False,
             demo_details: DemoCompanyDetails | None = None) -> FlowResult:
    """Compose existing contracts. `attest` is an external trusted boundary, never a default.

    Published P2 must be injected explicitly. The orchestrator never imports a test fixture.
    Invalid answers remain in the trace; only currently requested fields may be answered.
    """
    if demo_details is not None and not test_only:
        raise ValueError("Fictional company details require explicit TEST ONLY mode")
    profile, settings, catalog, prices, rules = deepcopy((profile, settings, catalog, prices, rules))
    extracted = extractor.extract(raw_text)
    state = deepcopy(extracted)
    initial = tuple(build_clarifications(state))
    updates = []
    for answer in answers:
        question = next((q for q in build_clarifications(state)
                         if q.field_name == answer.field_name and q.position == answer.position), None)
        if question is None:
            raise ValueError(f"Answer is not for a pending clarification: {answer.field_name}/{answer.position}")
        update = apply_clarification_answer(state, question, deepcopy(answer.value))
        updates.append(deepcopy(update))
        state = update.needs
    remaining = tuple(build_clarifications(state))
    selection = select_products(state, profile.products, product_ids=product_ids)
    # run_p2 itself validates/attempts prepare; a failed prepare must not be forced through.
    calculation = run_p2(state, selection, profile, settings=settings, catalog=catalog, prices=prices,
                         rules=rules, calculate_offer=calculate_offer, product_ids=product_ids)
    current = None
    if calculation.invocation is not None:
        current = prepare_p2_invocation(state, selection, profile, settings=settings, catalog=catalog,
                                        prices=prices, rules=rules, product_ids=product_ids)
    assessment = assess_selection(state, selection, profile, product_ids=product_ids,
                                  confirmed_sources=confirmed_sources, calculation_context=(
                                      current.bridge_context if current is not None else None),
                                  calculation=calculation.evidence_for(current) if current is not None else None,
                                  case_records=case_records, client_task=client_task, task_sources=task_sources)
    if attest is not None:
        if attestation is not None:
            raise ValueError("Supply either an attestation or an external attestor")
        attestation = attest(deepcopy(calculation))
    decision = evaluate_gate(calculation, current, assessment, attestation=attestation, as_of=as_of)
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    # Preserve previous review samples (including files currently open in Word).
    requested = path
    suffix = 1
    while path.exists():
        path = requested.with_name(f"{requested.stem}_{suffix}{requested.suffix}")
        suffix += 1
    if decision.status == GateStatus.COMMERCIALLY_APPROVED:
        path = generate_commercial_document(calculation, current, assessment, path,
                                            attestation=attestation, as_of=as_of, test_only=test_only,
                                            demo_details=demo_details)
    else:
        path = generate_internal_draft(decision, path)
    document = Document(path)
    if test_only and decision.status == GateStatus.INTERNAL_DRAFT:
        document.paragraphs[0].insert_paragraph_before(
            "ТЕСТОВЫЙ ОБРАЗЕЦ — НЕ ДЛЯ КОММЕРЧЕСКОГО ИСПОЛЬЗОВАНИЯ (TEST ONLY)")
    if decision.status == GateStatus.INTERNAL_DRAFT:
        document.add_heading("Уточнения и история ответов", 1)
        for value in (raw_text, initial, updates, remaining, extracted):
            document.add_paragraph(str(value))
    document.save(path)
    return FlowResult(raw_text, deepcopy(extracted), deepcopy(state), initial, tuple(updates), remaining,
                      selection, calculation, assessment, decision, path)

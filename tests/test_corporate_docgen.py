"""Corporate layout regressions through the real local flow; all positive attestations are TEST ONLY."""

from copy import deepcopy
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
import test_p2_published as published
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from test_docgen_gate import setup_gate
from test_orchestrator import _test_attestor, doc_text

from kp_plugin.docgen.builder import build_offer_document
from kp_plugin.docgen.corporate import HEADERS, CorporateOffer
from kp_plugin.docgen.gate import GateStatus
from kp_plugin.llm.cases import CommercialPermission
from kp_plugin.llm.needs import FakeLLMClient, NeedsExtractor
from kp_plugin.llm.p2_adapter import P2Settings
from kp_plugin.llm.p4_catalog import load_p4_catalog
from kp_plugin.orchestrator import Answer, run_flow

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'out/p3_e2e'
TODAY = date(2026, 10, 28)


def test_styled_test_only_offer_uses_bound_price_and_real_p2():
    data, _, _, assessment, _ = setup_gate(CommercialPermission.PROHIBITED)
    data['settings'] = data['settings'].model_copy(update={'p2_implementation': published.COMMIT})
    original = deepcopy(data)
    result = run_flow('TEST ONLY: нужна Учебный товар',
                      extractor=NeedsExtractor(FakeLLMClient({'items': [{'product_name': 'Учебный товар'}]})),
                      answers=(Answer('quantity', 1, '2.5'), Answer('unit', 1, 'т')),
                      profile=data['profile'], settings=data['settings'], catalog=data['catalog'],
                      prices=data['prices'], rules=data['rules'], calculate_offer=published.core.__wrapped__(),
                      output_path=OUTPUT/'styled_test_only_offer.docx', confirmed_sources=data['settings'].input_sources,
                      case_records=assessment.case_records, client_task='TEST TASK', task_sources=data['settings'].input_sources,
                      attest=_test_attestor, as_of=TODAY, test_only=True)
    assert result.decision.status == GateStatus.COMMERCIALLY_APPROVED
    document = Document(result.document)
    table = document.tables[0]
    assert len(table.columns) == 6
    assert tuple(cell.text for cell in table.rows[0].cells) == HEADERS
    assert table.rows[1].cells[2].text == '100,00 ₽'  # Input price, NOT 275/2.5.
    assert table.rows[1].cells[-1].text == '275,00 ₽'
    assert document.tables[1].rows[-1].cells[0].text == 'ОБЩАЯ СУММА С НДС'
    assert table.rows[0].cells[0]._tc.xml.find('163A59') >= 0
    assert table.rows[0].cells[0].paragraphs[0].runs[0].font.color.rgb.__str__() == 'FFFFFF'
    assert table.rows[1].cells[-1].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert 'EAF2F6' in document.tables[1].rows[-1].cells[0]._tc.xml
    assert table.autofit is False
    assert len(table.rows[0]._tr.xpath('./w:trPr/w:tblHeader')) == 1
    assert all(row._tr.xpath('./w:trPr/w:cantSplit') for row in table.rows)
    section = document.sections[0]
    assert abs(section.page_width.mm - 210) < 0.1 and abs(section.page_height.mm - 297) < 0.1
    assert all(18 <= margin.mm <= 22 for margin in (section.top_margin, section.bottom_margin,
                                                   section.left_margin, section.right_margin))
    assert abs(sum(column.width.mm for column in table.columns) - 170) < 0.2
    assert document.styles['Normal'].font.name == 'Arial'
    assert 'PAGE' in section.footer._element.xml
    text = doc_text(result.document)
    assert 'ТЕСТОВЫЙ ОБРАЗЕЦ — НЕ ДЛЯ КОММЕРЧЕСКОГО ИСПОЛЬЗОВАНИЯ' in text
    for forbidden in ('TEST ONLY CASE CONTENT', 'test-price', 'test-catalog', 'scope_complete', 'КП №',
                      'Контакты', 'Условия поставки', 'ИНН', 'Оптовые поставки рыбной продукции'):
        assert forbidden not in text
    assert data == original


def test_styled_real_p4_draft_has_no_fictitious_cost_table():
    base = Path(__file__).parent/'fixtures/p4_synthetic'
    paths = (base/'catalog/catalog.csv', base/'prices_stock/warehouse_stock.csv')
    original = tuple(path.read_bytes() for path in paths)
    profile = load_p4_catalog(*paths, as_of=TODAY)
    name = profile.products[0].name
    result = run_flow(f'Нужна {name}', extractor=NeedsExtractor(FakeLLMClient({'items': [{'product_name': name}]})),
                      profile=profile, settings=P2Settings(p2_implementation=published.COMMIT),
                      catalog={}, prices={}, rules={}, calculate_offer=published.core.__wrapped__(),
                      output_path=OUTPUT/'styled_internal_draft.docx', as_of=TODAY)
    assert result.decision.status == GateStatus.INTERNAL_DRAFT
    document = Document(result.document)
    assert not document.tables
    text = doc_text(result.document)
    assert 'ВНУТРЕННИЙ ЧЕРНОВИК — НЕ ДЛЯ КЛИЕНТА' in text
    assert 'INTERNAL_DRAFT' in text and 'Стоимость не подтверждена' in text
    assert result.remaining_questions
    assert 'synthetic' in text and 'quantity' in text and 'unit' in text
    assert result.calculation.raw_output is None
    assert 'ОБЩАЯ СУММА С НДС' not in text
    assert tuple(path.read_bytes() for path in paths) == original


def test_optional_business_fields_only_render_when_provided(tmp_path):
    offer = CorporateOffer(product='TEST ONLY товар', quantity=Decimal('2.5'), unit='т',
                           totals={'base_amount': Decimal('250.00'), 'discount_amount': Decimal('0.00'),
                                   'tax': Decimal('25.00'), 'total': Decimal('275.00')},
                           currency='RUB', document_date='2026-10-28', company_name='TEST ONLY Company',
                           client_name='TEST ONLY Client', offer_number='TEST-1', contacts=('TEST contact',),
                           delivery_terms=('TEST payment condition',), test_only=True)
    path = build_offer_document(offer, tmp_path/'optional.docx')
    document = Document(path)
    assert document.tables[0].rows[1].cells[2].text == 'Не подтверждена'
    assert 'TEST contact' in doc_text(path) and 'TEST payment condition' in doc_text(path)
    assert 'TEST ONLY Company' in document.sections[0].footer.paragraphs[0].text


@pytest.mark.parametrize('value', [Decimal('NaN'), Decimal('Infinity')])
def test_invalid_money_is_not_rendered_as_commercial_amount(value, tmp_path):
    offer = CorporateOffer(product='TEST', quantity=Decimal('2.5'), unit='т',
                           totals={'base_amount': value, 'discount_amount': Decimal('0.00'),
                                   'tax': Decimal('0.00'), 'total': value}, currency='RUB', document_date='2026-10-28')
    with pytest.raises(ValueError, match='finite original Decimal'):
        build_offer_document(offer, tmp_path/'invalid.docx')

"""Gate tests use explicitly labelled TEST ONLY attestations, never real commercial claims."""

from copy import deepcopy
from dataclasses import replace
from datetime import date
from decimal import Decimal

import pytest
from docx import Document
from test_p2_adapter import bundle, calculate_offer

from kp_plugin.docgen.gate import (
    CommercialAttestation,
    ConstraintReview,
    ConstraintStatus,
    GateStatus,
    constraint_review_fingerprint,
    evaluate_gate,
    generate_commercial_document,
    generate_internal_draft,
)
from kp_plugin.llm.cases import CaseRecord, CaseStatus, CommercialPermission
from kp_plugin.llm.p2_adapter import prepare_p2_invocation, run_p2
from kp_plugin.llm.p4_catalog import P4Row
from kp_plugin.llm.selection import select_products
from kp_plugin.llm.selection_bridge import assess_selection, canonical_fingerprint
from kp_plugin.llm.unreachable import ConfirmedSource

TODAY = date(2026, 10, 28)


def setup_gate(permission=None, content_confirmed=True):
    data = bundle()
    records = ()
    if permission is not None:
        source = data["profile"].products[0].sources[0]
        row = P4Row(1, (("product_id", "P1"), ("related_cases", "TEST-CASE"),
                        ("status", "active"), ("is_fictional", "нет")), source, ())
        data["profile"] = replace(data["profile"], catalog_rows=(row,))
        records = (CaseRecord(case_id="TEST-CASE", product_ids=("P1",), task_descriptions=("TEST TASK",),
                              content="TEST ONLY CASE CONTENT", sources=(source,), status=CaseStatus.VERIFIED,
                              source_confirmed=True, content_confirmed=content_confirmed, is_fictional=False,
                              commercial_permission=permission),)
    data["settings"] = data["settings"].model_copy(update={
        "delivery_cost": Decimal("0.0"), "delivery_sources": data["settings"].input_sources,
        "other_unknown_costs": (),
    })
    result = run_p2(**data, calculate_offer=calculate_offer)
    current = prepare_p2_invocation(**data)
    assessment = assess_selection(data["needs"], data["selection"], data["profile"],
                                  calculation=result.evidence, calculation_context=current.bridge_context,
                                  confirmed_sources=data["settings"].input_sources, case_records=records,
                                  client_task="TEST TASK", task_sources=data["settings"].input_sources)
    attestation = CommercialAttestation(
        issuer="TEST ONLY: not valid for real commercial use", invocation_fingerprint=current.fingerprint,
        output_fingerprint=canonical_fingerprint(result.raw_output), valid_until=TODAY,
        price_sources=(ConfirmedSource(source_id="test-price", version="v1", locator="P1", confirmed=True),),
        rule_sources=(ConfirmedSource(source_id="TEST ONLY rules", version="v1", locator="rules", confirmed=True),),
        scope_complete=True,
        constraints=ConstraintReview(status=ConstraintStatus.CONFIRMED,
                                     context_fingerprint=constraint_review_fingerprint(result, current),
                                     sources=data["settings"].input_sources, absence_confirmed=True,
                                     product_compatible=True, calculation_compatible=True, terms_compatible=True),
    )
    return data, result, current, assessment, attestation


def test_explicit_test_attestations_allow_checked_document(tmp_path):
    data, result, current, assessment, attestation = setup_gate()
    before = deepcopy((data, result, assessment))
    decision = evaluate_gate(result, current, assessment, attestation=attestation, as_of=TODAY)
    assert decision.status == GateStatus.COMMERCIALLY_APPROVED, decision.blockers
    path = generate_commercial_document(result, current, assessment, tmp_path / "offer.docx",
                                        attestation=attestation, as_of=TODAY)
    assert path.exists()
    assert (data, result, assessment) == before


@pytest.mark.parametrize("change", ["missing", "expired", "profile", "price", "rules", "quantity",
                                   "synthetic", "stock", "constraints", "delivery", "costs", "multi",
                                   "price_source", "rule_source", "output", "trust", "scope", "unit", "missing_version"])
def test_fail_closed_on_incomplete_or_changed_inputs(change, tmp_path):
    _data, result, current, assessment, attestation = setup_gate()
    if change == "missing":
        attestation = None
    elif change == "expired":
        attestation = replace(attestation, valid_until=date(2025, 1, 1))
    elif change in {"profile", "price"}:
        key = "profile_version" if change == "profile" else "price_snapshot_id"
        current = current.model_copy(update={"settings": current.settings.model_copy(update={key: "v2"})})
    elif change == "rules":
        current.rules["discounts"]["rules"][0]["percent"] = 10
    elif change == "quantity":
        assessment.needs.items[0].quantity = Decimal("3.0")
    elif change == "unit":
        assessment.needs.items[0].unit = "кг"
    elif change == "missing_version":
        current = current.model_copy(update={"settings": current.settings.model_copy(update={"profile_version": None})})
    elif change == "synthetic":
        current.catalog["P1"]["is_fictional"] = "да"
    elif change in {"stock", "constraints"}:
        product = assessment.profile.products[0].model_copy(update=(
            {"stock_quantity": None, "stock_sources": ()} if change == "stock" else
            {"constraints_check": "unknown", "constraint_sources": ()}))
        profile = replace(assessment.profile, products=(product,))
        assessment = replace(assessment, profile=profile,
                             selection=select_products(assessment.needs, profile.products))
    elif change in {"delivery", "costs"}:
        key, value = ("delivery_cost", None) if change == "delivery" else ("other_unknown_costs", ("insurance",))
        current = current.model_copy(update={"settings": current.settings.model_copy(update={key: value})})
    elif change == "multi":
        assessment.needs.items.append(deepcopy(assessment.needs.items[0]))
    elif change == "price_source":
        attestation = replace(attestation, price_sources=())
    elif change == "rule_source":
        attestation = replace(attestation, rule_sources=())
    elif change == "output":
        result.raw_output["total"] = Decimal("1.0")
    elif change == "trust":
        result = replace(result, trusted_inputs=False)
    else:
        attestation = replace(attestation, scope_complete=False)
    decision = evaluate_gate(result, current, assessment, attestation=attestation, as_of=TODAY)
    assert decision.status == GateStatus.INTERNAL_DRAFT
    assert decision.blockers
    destination = tmp_path / "blocked.docx"
    with pytest.raises(ValueError, match="Commercial generation blocked"):
        generate_commercial_document(result, current, assessment, destination,
                                     attestation=attestation, as_of=TODAY)
    assert not destination.exists()


def test_unknown_prices_draft_preserves_diagnostics(tmp_path):
    data, result, current, assessment, _ = setup_gate()
    result = replace(result, diagnostics=("TEST ONLY warning",), error_text="opaque error")
    decision = evaluate_gate(result, current, assessment, as_of=TODAY)
    path = generate_internal_draft(decision, tmp_path / "draft.docx")
    text = "\n".join(p.text for p in Document(path).paragraphs)
    assert "INTERNAL_DRAFT" in text and "TEST ONLY warning" in text and "opaque error" in text
    assert "commercial_attestation_missing" in text
    assert decision.calculation.needs == data["needs"]


@pytest.mark.parametrize("permission", [CommercialPermission.PROHIBITED, CommercialPermission.UNKNOWN])
def test_forbidden_or_unconfirmed_case_not_rendered(permission, tmp_path):
    _, result, current, assessment, attestation = setup_gate(permission)
    assert assessment.case_matches[0].case_id == "TEST-CASE"
    path = generate_commercial_document(result, current, assessment, tmp_path / "no-case.docx",
                                        attestation=attestation, as_of=TODAY)
    assert "TEST ONLY CASE CONTENT" not in "\n".join(p.text for p in Document(path).paragraphs)


@pytest.mark.parametrize("confirmed", [True, False])
def test_only_confirmed_allowed_case_content_can_be_rendered(confirmed, tmp_path):
    _, result, current, assessment, attestation = setup_gate(CommercialPermission.ALLOWED, confirmed)
    path = generate_commercial_document(result, current, assessment, tmp_path / "case.docx",
                                        attestation=attestation, as_of=TODAY)
    text = "\n".join(p.text for p in Document(path).paragraphs)
    assert ("TEST ONLY CASE CONTENT" in text) is confirmed

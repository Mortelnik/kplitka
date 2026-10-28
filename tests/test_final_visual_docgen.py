"""Final visual content contract; fictional company information is confined to TEST ONLY."""

import json
from dataclasses import replace
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
import test_p2_published as published
from docx import Document
from test_docgen_gate import setup_gate
from test_orchestrator import _test_attestor, doc_text

from kp_plugin.docgen.builder import build_offer_document
from kp_plugin.docgen.corporate import HEADERS, CorporateOffer, DemoCompanyDetails
from kp_plugin.docgen.gate import (
    GateStatus,
    evaluate_gate,
    generate_commercial_document,
    generate_internal_draft,
)
from kp_plugin.llm.needs import FakeLLMClient, NeedsExtractor
from kp_plugin.orchestrator import Answer, run_flow

ROOT = Path(__file__).resolve().parents[1]
TODAY = date(2026, 10, 28)


def demo_details():
    data = json.loads((ROOT/'tests/fixtures/demo_company_test_only.json').read_text(encoding='utf-8'))
    assert data.pop('test_only') is True
    for field in ('contacts', 'delivery_terms'):
        data[field] = tuple(data[field])
    return DemoCompanyDetails(**data)


def test_final_visual_offer_runs_real_p2_and_displays_balanced_test_content():
    data, *_ = setup_gate()
    data['settings'] = data['settings'].model_copy(update={'p2_implementation': published.COMMIT})
    flow = run_flow('TEST ONLY нужна Учебный товар',
                    extractor=NeedsExtractor(FakeLLMClient({'items': [{'product_name': 'Учебный товар'}]})),
                    answers=(Answer('quantity', 1, '2.5'), Answer('unit', 1, 'т')),
                    profile=data['profile'], settings=data['settings'], catalog=data['catalog'], prices=data['prices'],
                    rules=data['rules'], calculate_offer=published.core.__wrapped__(),
                    output_path=ROOT/'out/p3_e2e/final_visual_offer.docx',
                    confirmed_sources=data['settings'].input_sources, attest=_test_attestor, as_of=TODAY,
                    test_only=True, demo_details=demo_details())
    assert flow.decision.status == GateStatus.COMMERCIALLY_APPROVED
    assert flow.calculation.raw_output['total'] == Decimal('275.00')
    assert flow.selection.positions[0].candidates[0].product_id == 'P1'
    document = Document(flow.document)
    assert document.paragraphs[0].text.startswith('ТЕСТОВЫЙ ОБРАЗЕЦ — НЕ ДЛЯ КОММЕРЧЕСКОГО ИСПОЛЬЗОВАНИЯ')
    assert document.paragraphs[0].runs[0].bold
    assert 'EAF2F6' in document.paragraphs[0]._p.xml
    text = doc_text(flow.document)
    assert '28 октября 2026 г.' in text
    assert 'Валюта: RUB' not in text and '2026-10-28' not in text
    assert 'Товар P1:' not in text and 'P1' not in text
    assert tuple(c.text for c in document.tables[0].rows[0].cells) == HEADERS
    assert document.tables[0].rows[1].cells[2].text == '100,00 ₽'
    assert document.tables[0].rows[1].cells[-1].text == '275,00 ₽'
    assert '163A59' in document.tables[0].rows[0].cells[0]._tc.xml
    assert 'EAF2F6' in document.tables[1].rows[-1].cells[0]._tc.xml
    for value in ('Северная линия', 'Ресторан у моря', 'Условия поставки', 'Контакты', 'Менеджер:', 'Телефон:',
                  'manager@example.invalid', 'Оплата:', 'Срок:', 'Способ:'):
        assert value in text
    assert 'source_ref' not in text and 'КП №' not in text
    assert not document.element.body.xpath('.//w:br[@w:type="page"]')


def test_fictional_details_rejected_by_normal_gate_path(tmp_path):
    _, result, current, assessment, attestation = setup_gate()
    with pytest.raises(ValueError, match='explicit TEST ONLY'):
        generate_commercial_document(result, current, assessment, tmp_path/'forbidden.docx',
                                     attestation=attestation, as_of=TODAY, demo_details=demo_details())
    assert not (tmp_path/'forbidden.docx').exists()
    path = generate_commercial_document(result, current, assessment, tmp_path/'normal.docx',
                                        attestation=attestation, as_of=TODAY)
    text = doc_text(path)
    for absent in ('Северная линия', 'Ресторан у моря', 'manager@example.invalid', 'Условия поставки', 'Контакты',
                   'Клиент:', 'Товар P1:', 'Валюта: RUB'):
        assert absent not in text
    assert '28 октября 2026 г.' in text and '88,00 ₽' in text  # Existing gate protocol fixture, not the real-P2 example.


def test_explicit_test_only_does_not_bypass_gate(tmp_path):
    _, result, current, assessment, _ = setup_gate()
    with pytest.raises(ValueError, match='Commercial generation blocked'):
        generate_commercial_document(result, current, assessment, tmp_path/'no-attestation.docx',
                                     test_only=True, demo_details=demo_details(), as_of=TODAY)


def test_draft_remains_marked_and_without_demo_content(tmp_path):
    _, result, current, assessment, _ = setup_gate()
    decision = evaluate_gate(result, current, assessment, as_of=TODAY)
    path = generate_internal_draft(decision, tmp_path/'draft.docx')
    document = Document(path)
    assert document.paragraphs[0].text == 'ВНУТРЕННИЙ ЧЕРНОВИК — НЕ ДЛЯ КЛИЕНТА'
    assert not document.tables
    assert 'Северная линия' not in doc_text(path)


def test_long_product_and_amounts_keep_explicit_widths_and_all_values(tmp_path):
    value = Decimal('123456789012345.67')
    offer = CorporateOffer(product='Очень длинное название рыбной продукции ' * 12, quantity=Decimal('123456.789'),
                           unit='т', totals={'base_amount': value, 'discount_amount': Decimal('0.00'),
                                            'tax': Decimal('0.00'), 'total': value},
                           currency='RUB', document_date='2026-10-28')
    path = build_offer_document(offer, tmp_path/'long.docx')
    document = Document(path)
    table = document.tables[0]
    assert len(table.columns) == 6 and table.autofit is False
    assert table.rows[1].cells[0].text == offer.product
    assert '123 456 789 012 345,67 ₽' == table.rows[1].cells[-1].text
    assert abs(sum(c.width.mm for c in table.columns) - 170) < 0.2
    assert all(row._tr.xpath('./w:trPr/w:cantSplit') for row in table.rows)
    assert not table._tbl.xpath('.//w:noWrap')
    assert table.rows[1].height is None


def test_provided_real_presentation_values_and_empty_optional_fields(tmp_path):
    # Explicit renderer input, not an assertion of real-world commercial truth.
    offer = CorporateOffer(product='TEST товар', quantity=Decimal('1.0'), unit='т',
                           totals={key: Decimal('0.00') for key in ('base_amount', 'discount_amount', 'tax', 'total')},
                           currency='RUB', document_date='2026-10-28', company_name='TEST provided company',
                           business_description='TEST provided activity', client_name='TEST provided client')
    path = build_offer_document(offer, tmp_path/'provided.docx')
    assert 'TEST provided company' in doc_text(path)
    assert 'TEST provided activity' in doc_text(path)
    path = build_offer_document(replace(offer, company_name=None, business_description=None, client_name=None),
                                tmp_path/'empty.docx')
    assert 'TEST provided' not in doc_text(path)

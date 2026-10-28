"""Reproducible local text-to-DOCX examples; FakeLLM returns a specified payload, not language understanding."""

from copy import deepcopy
from dataclasses import replace
from datetime import date
from pathlib import Path

import pytest
import test_p2_published as published
from docx import Document
from test_docgen_gate import setup_gate

from kp_plugin.docgen.gate import (
    CommercialAttestation,
    ConstraintReview,
    ConstraintStatus,
    GateStatus,
    constraint_review_fingerprint,
)
from kp_plugin.llm.cases import CommercialPermission
from kp_plugin.llm.needs import FakeLLMClient, NeedsExtractor
from kp_plugin.llm.p2_adapter import P2Settings
from kp_plugin.llm.p4_catalog import load_p4_catalog
from kp_plugin.llm.selection_bridge import canonical_fingerprint
from kp_plugin.llm.unreachable import ConfirmedSource
from kp_plugin.orchestrator import Answer, run_flow

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "out" / "p3_e2e"
TODAY = date(2026, 10, 28)
COMMIT = published.COMMIT


@pytest.fixture
def core():
    return published.core.__wrapped__()


def doc_text(path):
    document = Document(path)
    return "\n".join([p.text for p in document.paragraphs]
                     + [cell.text for table in document.tables for row in table.rows for cell in row.cells])


def test_real_p4_csv_to_internal_draft_without_forcing_p2(core):
    base = Path(__file__).parent / "fixtures/p4_synthetic"
    catalog_path, stock_path = base / "catalog/catalog.csv", base / "prices_stock/warehouse_stock.csv"
    before = catalog_path.read_bytes(), stock_path.read_bytes()
    profile = load_p4_catalog(catalog_path, stock_path, as_of=TODAY)
    fingerprint = canonical_fingerprint(profile)
    product = profile.products[0]
    calls = []

    def observed(**kwargs):
        calls.append(kwargs)
        return core(**kwargs)

    flow = run_flow(f"Здравствуйте, нужна {product.name}",
                    extractor=NeedsExtractor(FakeLLMClient({"items": [{"product_name": product.name}]})),
                    answers=(Answer("quantity", 1, "2.5"), Answer("unit", 1, "т")), profile=profile,
                    settings=P2Settings(profile_id="P4 educational", profile_version=profile.catalog_rows[0].source.version,
                                        p2_implementation=COMMIT), catalog={}, prices={}, rules={}, calculate_offer=observed,
                    output_path=OUTPUT / "p4_internal_draft.docx", as_of=TODAY)
    assert not calls
    assert flow.calculation.raw_output is None and flow.calculation.evidence is None
    assert flow.decision.status == GateStatus.INTERNAL_DRAFT
    assert "product_identity_source_unconfirmed" in flow.calculation.diagnostics
    assert flow.assessment.profile.stock_rows and all(row.issues for row in flow.assessment.profile.stock_rows)
    assert flow.initial_questions and not flow.remaining_questions
    assert flow.extracted.items[0].quantity is None
    assert flow.needs.items[0].quantity is not None
    assert flow.extracted.items[0].evidence["product_name"] == flow.needs.items[0].evidence["product_name"]
    assert flow.needs.items[0].evidence["quantity"].verified
    text = doc_text(flow.document)
    for marker in ("INTERNAL_DRAFT", "НЕ ДЛЯ КЛИЕНТА", "synthetic", "stock_quantity", "constraints", "quantity"):
        assert marker in text
    assert flow.decision.assessment.needs.model_payload == flow.needs.model_payload
    assert canonical_fingerprint(profile) == fingerprint
    assert (catalog_path.read_bytes(), stock_path.read_bytes()) == before


def _test_attestor(result):
    """Explicit TEST ONLY trusted caller, not a real authority or commercial confirmation."""
    return CommercialAttestation(
        issuer="TEST ONLY local simulation", invocation_fingerprint=result.invocation.fingerprint,
        output_fingerprint=canonical_fingerprint(result.raw_output), valid_until=TODAY,
        price_sources=(ConfirmedSource(source_id="test-price", version="v1", locator="P1", confirmed=True),),
        rule_sources=(ConfirmedSource(source_id="TEST ONLY rules", version="v1", locator="rules", confirmed=True),),
        scope_complete=True,
        constraints=ConstraintReview(status=ConstraintStatus.CONFIRMED,
                                     context_fingerprint=constraint_review_fingerprint(result, result.invocation),
                                     sources=result.invocation.settings.input_sources, absence_confirmed=True,
                                     product_compatible=True, calculation_compatible=True, terms_compatible=True),
    )


@pytest.mark.parametrize("permission", [CommercialPermission.PROHIBITED, CommercialPermission.UNKNOWN])
def test_test_only_text_to_actual_p2_and_gated_docx(core, permission):
    data, _, _, assessment, _ = setup_gate(permission)
    data["settings"] = data["settings"].model_copy(update={"p2_implementation": COMMIT})
    snapshot = canonical_fingerprint(data)
    text = "TEST ONLY: нужна Учебный товар"
    payload = {"items": [{"product_name": "Учебный товар", "quantity": "99", "unit": "кг"}]}
    original_payload = deepcopy(payload)
    path = OUTPUT / ("test_only_offer.docx" if permission == CommercialPermission.PROHIBITED else "test_only_unknown_case.docx")
    flow = run_flow(text, extractor=NeedsExtractor(FakeLLMClient(payload)),
                    answers=(Answer("quantity", 1, "2.5"), Answer("unit", 1, "т")),
                    profile=data["profile"], settings=data["settings"], catalog=data["catalog"], prices=data["prices"],
                    rules=data["rules"], calculate_offer=core, output_path=path,
                    confirmed_sources=data["settings"].input_sources, case_records=assessment.case_records,
                    client_task="TEST TASK", task_sources=data["settings"].input_sources,
                    attest=_test_attestor, as_of=TODAY, test_only=True)
    assert flow.decision.status == GateStatus.COMMERCIALLY_APPROVED, flow.decision.blockers
    assert flow.calculation.raw_output["total"] == 275
    assert flow.assessment.calculation_confirmed
    assert flow.calculation.invocation.settings.p2_implementation == COMMIT
    assert flow.calculation.invocation.base_context.price_snapshot_id == "prices-v1"
    assert flow.selection.positions[0].candidates[0].product_id == "P1"
    assert flow.extracted.items[0].unknown_fields == ["quantity", "unit"]
    assert not flow.needs.items[0].unknown_fields
    assert flow.needs.items[0].model_payload["quantity"] == "99"
    assert flow.needs.items[0].evidence_history["quantity"][0].verified is False
    assert payload == original_payload and canonical_fingerprint(data) == snapshot
    rendered = doc_text(flow.document)
    assert "TEST ONLY" in rendered and "НЕ ДЛЯ КОММЕРЧЕСКОГО" in rendered
    assert "275,00" in rendered
    assert "Товар P1:" not in rendered
    assert "TEST ONLY CASE CONTENT" not in rendered
    assert flow.assessment.case_matches[0].case_id == "TEST-CASE"
    assert not flow.decision.commercial_cases


def test_unanswered_and_invalid_answers_remain_draft(core, tmp_path):
    data, *_ = setup_gate()
    flow = run_flow("нужна Учебный товар", extractor=NeedsExtractor(FakeLLMClient({"items": [
        {"product_name": "Учебный товар"}]})), answers=(Answer("quantity", 1, "много"),),
        profile=data["profile"], settings=data["settings"], catalog=data["catalog"], prices=data["prices"],
        rules=data["rules"], calculate_offer=core, output_path=tmp_path / "pending.docx")
    assert flow.updates[0].status == "invalid"
    assert flow.remaining_questions and flow.decision.status == GateStatus.INTERNAL_DRAFT
    assert "много" not in str(flow.needs.items[0].quantity)
    assert "INTERNAL_DRAFT" in doc_text(flow.document)


def test_changed_price_rejects_external_attestation(core, tmp_path):
    data, _, _, _, old = setup_gate()
    flow = run_flow("Учебный товар 2.5 т", extractor=NeedsExtractor(FakeLLMClient({"items": [
        {"product_name": "Учебный товар", "quantity": "2.5", "unit": "т"}]})),
        profile=data["profile"], settings=data["settings"].model_copy(update={"price_snapshot_id": "v2"}),
        catalog=data["catalog"], prices=data["prices"], rules=data["rules"], calculate_offer=core,
        output_path=tmp_path / "stale.docx", confirmed_sources=data["settings"].input_sources,
        attestation=replace(old, valid_until=TODAY), as_of=TODAY, test_only=True)
    assert flow.decision.status == GateStatus.INTERNAL_DRAFT
    assert "invalid_or_expired_attestation" in flow.decision.blockers

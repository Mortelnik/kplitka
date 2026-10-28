"""Cosmetic changes only: typography/spacing and isolated TEST ONLY captions."""

from datetime import date
from decimal import Decimal
from pathlib import Path

import test_p2_published as published
from docx import Document
from test_docgen_gate import setup_gate
from test_final_visual_docgen import demo_details
from test_orchestrator import _test_attestor, doc_text

from kp_plugin.docgen.gate import GateStatus, evaluate_gate, generate_internal_draft
from kp_plugin.llm.needs import FakeLLMClient, NeedsExtractor
from kp_plugin.orchestrator import Answer, run_flow

ROOT = Path(__file__).resolve().parents[1]
TODAY = date(2026, 10, 28)


def test_polished_sample_fonts_spacing_and_unchanged_calculation(tmp_path):
    data, *_ = setup_gate()
    data['settings'] = data['settings'].model_copy(update={'p2_implementation': published.COMMIT})
    def create_document(path):
        return run_flow('TEST ONLY нужна Учебный товар',
                    extractor=NeedsExtractor(FakeLLMClient({'items': [{'product_name': 'Учебный товар'}]})),
                    answers=(Answer('quantity', 1, '2.5'), Answer('unit', 1, 'т')),
                    profile=data['profile'], settings=data['settings'], catalog=data['catalog'], prices=data['prices'],
                    rules=data['rules'], calculate_offer=published.core.__wrapped__(),
                    output_path=path,
                    confirmed_sources=data['settings'].input_sources, attest=_test_attestor, as_of=TODAY,
                    test_only=True, demo_details=demo_details())
    previous = create_document(tmp_path/'previous.docx').document
    before = previous.read_bytes()
    flow = create_document(tmp_path/'polished.docx')
    document = Document(flow.document)
    assert flow.decision.status == GateStatus.COMMERCIALLY_APPROVED
    assert document.styles['Normal'].font.name == 'Arial'
    assert document.styles['Normal'].font.size.pt == 11
    assert document.styles['Title'].font.size.pt == 23
    assert document.styles['Company Name'].font.size.pt == 14
    assert document.styles['Company Name'].font.bold
    assert str(document.styles['Company Name'].font.color.rgb) == '163A59'
    assert document.styles['Heading 1'].font.size.pt == 12
    assert document.styles['Heading 1'].font.bold
    company = next(p for p in document.paragraphs if p.text == 'ДЕМО «Северная линия»')
    assert company.style.name == 'Company Name'
    headings = {p.text: p for p in document.paragraphs}
    assert headings['Условия поставки'].paragraph_format.space_before.pt == 19
    assert headings['Контакты'].paragraph_format.space_before.pt == 19
    assert document.styles['Heading 1'].paragraph_format.space_before.pt == 14
    assert len(document.tables[0].columns) == 6
    for row in document.tables[0].rows:
        for cell in row.cells:
            assert all(9 <= r.font.size.pt <= 10 for r in cell.paragraphs[0].runs)
    assert document.tables[0].rows[1].cells[2].text == '100,00 ₽'
    assert document.tables[0].rows[1].cells[-1].text == '275,00 ₽'
    assert flow.calculation.raw_output['base_amount'] == Decimal('250.00')
    assert flow.calculation.raw_output['discount_amount'] == Decimal('0.00')
    assert flow.calculation.raw_output['tax'] == Decimal('25.00')
    assert flow.calculation.raw_output['total'] == Decimal('275.00')
    text = doc_text(flow.document)
    assert document.paragraphs[0].text.startswith('ТЕСТОВЫЙ ОБРАЗЕЦ — НЕ ДЛЯ КОММЕРЧЕСКОГО ИСПОЛЬЗОВАНИЯ')
    assert 'TEST ONLY' in text and 'ДЕМО' in text
    assert 'вымышленная компания' not in text and '(вымышленное лицо)' not in text
    assert 'manager@example.invalid' in text
    assert not document.element.body.xpath('.//w:br[@w:type="page"]')
    assert not document.element.body.xpath('.//w:pageBreakBefore')
    assert len(document.sections) == 1
    assert previous.read_bytes() == before


def test_internal_draft_marker_and_gate_are_unchanged(tmp_path):
    _, result, current, assessment, _ = setup_gate()
    decision = evaluate_gate(result, current, assessment, as_of=TODAY)
    assert decision.status == GateStatus.INTERNAL_DRAFT
    path = generate_internal_draft(decision, tmp_path/'draft.docx')
    document = Document(path)
    assert document.paragraphs[0].text == 'ВНУТРЕННИЙ ЧЕРНОВИК — НЕ ДЛЯ КЛИЕНТА'
    assert 'Северная линия' not in doc_text(path)

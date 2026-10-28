"""Compact review remains internal; all inputs are explicitly TEST ONLY fixtures."""

from datetime import date

import pytest
import test_p2_published as published
from docx import Document
from test_docgen_gate import setup_gate

from kp_plugin.docgen.gate import GateStatus, generate_priced_internal_draft
from kp_plugin.llm.needs import FakeLLMClient, NeedsExtractor
from kp_plugin.orchestrator import Answer, run_flow


def make_review(tmp_path):
    data, *_ = setup_gate()
    settings = data["settings"].model_copy(update={"p2_implementation": published.COMMIT})
    return run_flow(
        "TEST ONLY: " + data["profile"].products[0].name,
        extractor=NeedsExtractor(FakeLLMClient({"items": [{"product_name": data["profile"].products[0].name}]})),
        answers=(Answer("quantity", 1, "2.5"), Answer("unit", 1, "т")),
        profile=data["profile"], settings=settings, catalog=data["catalog"], prices=data["prices"],
        rules=data["rules"], calculate_offer=published.core.__wrapped__(),
        output_path=tmp_path / "diagnostics.docx", confirmed_sources=settings.input_sources,
        as_of=date(2026, 10, 28), test_only=True,
    )


def test_compact_review_retains_internal_status_and_original_amounts(tmp_path):
    flow = make_review(tmp_path)
    assert flow.decision.status == GateStatus.INTERNAL_DRAFT
    assert "commercial_attestation_missing" in flow.decision.blockers
    path = generate_priced_internal_draft(flow.decision, tmp_path / "review.docx", as_of=date(2026, 10, 28))
    doc = Document(path)
    body = "\n".join(p.text for p in doc.paragraphs)
    table_text = " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "INTERNAL_DRAFT / TEST ONLY" in body
    assert "TEST ONLY" in doc.sections[0].footer.paragraphs[-1].text
    for unwanted in ("ДЕМО", "Северная линия", "Ресторан у моря", "manager@example.invalid", "учебный самовывоз"):
        assert unwanted not in body
    assert body.count("Уточняется при согласовании заказа") == 3
    assert "Контактные данные не предоставлены" in body
    assert len(doc.tables[0].columns) == 6
    for amount in ("100,00", "0,00", "25,00", "275,00"):
        assert amount in table_text
    assert flow.calculation.raw_output["total"] == 275
    assert not doc.element.xpath(".//w:br[@w:type='page']")


def test_priced_review_rejects_commercial_decision(tmp_path):
    from dataclasses import replace

    flow = make_review(tmp_path)
    with pytest.raises(ValueError, match="INTERNAL_DRAFT"):
        generate_priced_internal_draft(replace(flow.decision, status=GateStatus.COMMERCIALLY_APPROVED),
                                      tmp_path / "rejected.docx", as_of=date(2026, 10, 28))
